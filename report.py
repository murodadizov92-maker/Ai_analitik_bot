"""Agentlarning kunlik hisoboti: Sales Doctor ma'lumotlarini yig'ish, hisoblash, matn va Excel."""
import datetime as dt
import html
import io
from collections import defaultdict


# Maydonda ishlamaydigan 'agentlar' (ofis, to'g'ridan-to'g'ri savdo): vizit KPI va agent tahlilidan chiqariladi
NON_FIELD = {"ofis"}
IGNORE = set()   # ishdan ketgan agentlar (normallashtirilgan ismlar): ro'yxatlarda ko'rsatilmaydi

SHIP_FILTER = "dateLoad"      # "yuk chiqqan" sanasi: getOrder filtri (hujjat/otgruzka sanasi)
SHIP_STATUSES = [2, 3, 4]    # yuborilgan, yetkazilgan, yopilgan
VISIT_CAP_MIN = 90      # bundan uzoq yozilgan vizit yopilmay qolgan hisoblanadi
SHORT_VISIT_SEC = 60    # bundan qisqa vizit shubhali hisoblanadi


# ---------------------------------------------------------------- yordamchilar
def norm_name(x):
    """Ismlarni solishtirish uchun: kichik harf, ortiqcha probellar va qavs ichidagi probellar olib tashlanadi."""
    import re
    x = re.sub(r"\s+", " ", (x or "").strip().lower())
    x = re.sub(r"\(\s+", "(", x)
    return re.sub(r"\s+\)", ")", x)


def money(x):
    return f"{x:,.0f}".replace(",", " ")


def short(x):
    x = float(x)
    if abs(x) >= 1e9:
        return f"{x / 1e9:.2f} mlrd"
    if abs(x) >= 1e6:
        return f"{x / 1e6:.0f} mln"
    return money(x)


def status(pct):
    return "🟢" if pct >= 1.0 else ("🟡" if pct >= 0.85 else "🔴")


def pdt(s):
    if not s:
        return None
    for chunk, fmt in ((s[:19], "%Y-%m-%d %H:%M:%S"), (s[:16], "%Y-%m-%d %H:%M")):
        try:
            return dt.datetime.strptime(chunk, fmt)
        except ValueError:
            pass
    return None


def fmt_min(m):
    m = int(round(m))
    h, mm = divmod(m, 60)
    return f"{h} soat {mm} daq" if h else f"{mm} daq"


def sid(obj):
    return (obj or {}).get("SD_id")


# ---------------------------------------------------------------- ma'lumot yig'ish
async def collect(client, d_from, d_to, statuses):
    """d_from..d_to oralig'i bo'yicha (bir kun uchun d_from == d_to)."""
    per = {"period": {"date": {"from": str(d_from), "to": str(d_to)}}}
    agents = await client.paginate("getAgent", {}, "agent")
    visits = await client.paginate("getVisit", {"filter": per}, "visit")
    orders = await client.paginate("getOrder", {"filter": {"agent": "all", "status": statuses, **per}}, "order")
    shipped = await client.paginate(
        "getOrder", {"filter": {"agent": "all", "status": SHIP_STATUSES,
                                "period": {SHIP_FILTER: {"from": str(d_from), "to": str(d_to)}}}}, "order")
    pays = await client.paginate("getPayment", {"filter": {"transactionType": 3, **per}}, "payment")
    defects = await client.paginate("getOrderDefect", {"filter": {"status": statuses, **per}}, "order")
    ptypes = await client.paginate("getPaymentType", {}, "currency")

    cnames, pnames = {}, {}
    for v in visits:
        if v.get("client_id"):
            cnames[v["client_id"]] = v.get("client_name") or ""
    for o in orders:
        c = o.get("client") or {}
        if c.get("SD_id"):
            cnames[c["SD_id"]] = c.get("clientName") or cnames.get(c["SD_id"], "")
        for line in o.get("orderProducts") or []:
            p = line.get("product") or {}
            if p.get("SD_id") and p.get("name"):
                pnames[p["SD_id"]] = p["name"]

    # Nomi hali noma'lum mijoz/tovarlar bo'lsagina ro'yxatni to'liq o'qiymiz
    need_c = {sid(x.get("client")) for x in pays + defects} - {None}
    if any(not cnames.get(c) for c in need_c):
        for c in await client.paginate("getClient", {}, "client"):
            if not cnames.get(c["SD_id"]):
                cnames[c["SD_id"]] = c.get("name") or c["SD_id"]
    need_p = {sid(l.get("product")) for x in defects for l in x.get("defectProducts") or []} - {None}
    if need_p - set(pnames):
        for p in await client.paginate("getProduct", {}, "product"):
            pnames.setdefault(p["SD_id"], p.get("name") or p["SD_id"])

    return {"day": d_to, "d_from": d_from, "agents": agents, "visits": visits, "orders": orders, "shipped": shipped, "pays": pays,
            "defects": defects, "ptypes": ptypes, "cnames": cnames, "pnames": pnames}


# ---------------------------------------------------------------- hisoblash
def new_stat(name, aid=None):
    return {"ship_n": 0, "ship_sum": 0.0, "days": set(), "field": (name or "").strip().lower() not in NON_FIELD, "short": 0, "open_long": 0, "id": aid, "with_order": 0, "order_clients": set(), "lines": 0, "no_order_clients": [], "name": name, "planned": 0, "visited": 0, "gps": 0, "photo": 0, "no_photo": [], "no_gps": [],
            "first": None, "last": None, "in_visit": 0.0, "durs": [], "iv": [], "rejects": [],
            "orders_n": 0, "orders_sum": 0.0, "pay_sum": 0.0, "pay_by": defaultdict(float), "pays": [],
            "ret_n": 0, "ret_sum": 0.0, "rets": []}


def summarize(data):
    names = {a["SD_id"]: a.get("name") or a["SD_id"] for a in data["agents"]}
    ptn = {p["SD_id"]: p.get("name") or p["SD_id"] for p in data["ptypes"]}
    cn = lambda cid: (data["cnames"].get(cid) or cid or "—")
    pn = lambda pid: (data["pnames"].get(pid) or pid or "—")
    S = {}

    def st(aid, fallback=None):
        if aid not in S:
            S[aid] = new_stat(names.get(aid) or fallback or str(aid), aid)
        return S[aid]

    for v in data["visits"]:
        s = st(v.get("agent_id"), v.get("agent_name"))
        s["planned"] += int(v.get("planned") or 0)
        client = v.get("client_name") or v.get("client_id")
        if v.get("reject"):
            s["rejects"].append((client, v["reject"]))
        if int(v.get("visited") or 0):
            s["visited"] += 1
            s["gps"] += int(v.get("gps_visit") or 0)
            s["photo"] += int(v.get("has_photo_report") or 0)
            ho = int(v.get("has_order") or 0)
            s["with_order"] += ho
            if not ho and not v.get("reject"):
                s["no_order_clients"].append(client)
            if not int(v.get("has_photo_report") or 0):
                s["no_photo"].append(client)
            if not int(v.get("gps_visit") or 0):
                s["no_gps"].append(client)
            a, b = pdt(v.get("start_date")), pdt(v.get("end_date"))
            if a:
                s["days"].add(a.date())
                s["first"] = a if not s["first"] else min(s["first"], a)
                end = a
                if b and b > a:
                    d = (b - a).total_seconds()
                    if d < SHORT_VISIT_SEC:
                        s["short"] += 1
                    if d > VISIT_CAP_MIN * 60:
                        s["open_long"] += 1
                        d = VISIT_CAP_MIN * 60
                    end = a + dt.timedelta(seconds=d)
                    s["durs"].append(d / 60)
                s["last"] = end if not s["last"] else max(s["last"], end)
                s["iv"].append((a, end))
    for s in S.values():
        merged = []
        for a_, b_ in sorted(s["iv"]):
            if merged and a_ <= merged[-1][1]:
                merged[-1][1] = max(merged[-1][1], b_)
            else:
                merged.append([a_, b_])
        s["in_visit"] = sum((b_ - a_).total_seconds() for a_, b_ in merged) / 60
        gaps = [(merged[i + 1][0] - merged[i][1]).total_seconds() / 60 for i in range(len(merged) - 1)]
        s["max_gap"] = max(gaps or [0])

    for o in data["orders"]:
        s = st(sid(o.get("agent")))
        s["orders_n"] += 1
        s["order_clients"].add(sid(o.get("client")))
        s["lines"] += len(o.get("orderProducts") or [])
        s["orders_sum"] += float(o.get("totalSummaAfterDiscount") or o.get("totalSumma") or 0)

    for o in data.get("shipped", []):
        s = st(sid(o.get("agent")))
        s["ship_n"] += 1
        s["ship_sum"] += float(o.get("totalSummaAfterDiscount") or o.get("totalSumma") or 0)

    for p in data["pays"]:
        s = st(sid(p.get("agent")), "Agentsiz (kassa)")
        amt = float(p.get("amount") or 0)
        s["pay_sum"] += amt
        t = ptn.get(sid(p.get("paymentType")), "boshqa")
        s["pay_by"][t] += amt
        s["pays"].append((cn(sid(p.get("client"))), amt, t))

    for d in data["defects"]:
        s = st(sid(d.get("agent")), "Agentsiz")
        summa = float(d.get("summa") or d.get("totalSumma") or 0)
        s["ret_n"] += 1
        s["ret_sum"] += summa
        prods = [(pn(sid(l.get("product"))), l.get("quantity") or 0) for l in d.get("defectProducts") or []]
        s["rets"].append((cn(sid(d.get("client"))), summa, d.get("reason") or "", prods))

    # faol, lekin bugun hech narsa qilmagan agentlar
    idle = [a.get("name") or a["SD_id"] for a in data["agents"]
            if a.get("active", "Y") != "N" and a["SD_id"] not in S and norm_name(a.get("name")) not in IGNORE]
    return S, idle


# ---------------------------------------------------------------- Telegram matni
def warnings(s):
    w = []
    if s["field"] and s["planned"] and s["visited"] / s["planned"] < 0.6:
        w.append(f"reja bajarilishi {s['visited'] * 100 // s['planned']}%")
    if not s["field"]:
        return w
    if s["no_photo"]:
        w.append(f"{len(s['no_photo'])} ta vizitda foto otchyot yo'q")
    if s["no_gps"]:
        w.append(f"{len(s['no_gps'])} ta vizit GPS bilan tasdiqlanmagan")
    if s["short"] >= 3:
        w.append(f"{s['short']} ta vizit 1 daqiqadan qisqa (haqiqiyligi shubhali)")
    if s["open_long"]:
        w.append(f"{s['open_long']} ta vizit yopilmay qolgan ({VISIT_CAP_MIN}+ daq)")
    if s["field"] and not s["visited"] and (s["orders_n"] or s["pay_sum"]):
        w.append("vizitsiz buyurtma/to'lov bor")
    return w


def agent_block(s, limit=5, mon=None, multi=False):
    e = lambda x: html.escape(str(x), quote=False)
    L = [f"{'👤' if s['field'] else '🏢'} <b>{e(s['name'])}</b>" + ("" if s["field"] else " (filiallarga yuk / ofis savdosi)")]
    plan = f" / reja {s['planned']}" if s["planned"] else ""
    if s["field"] or s["visited"]:
        L.append(f"🚶 Vizit: <b>{s['visited']}</b>{plan} · GPS: {s['gps']} · Foto: {s['photo']}")
    if multi and s["field"] and s["days"]:
        L.append(f"📆 Ish kunlari: {len(s['days'])} · kuniga o'rtacha {s['visited'] / len(s['days']):.0f} vizit · "
                 f"vizitlarda jami {fmt_min(s['in_visit'])}")
    if s["first"] and s["field"] and not multi:
        total = (s["last"] - s["first"]).total_seconds() / 60
        avg = f" · o'rtacha vizit {fmt_min(sum(s['durs']) / len(s['durs']))}" if s["durs"] else ""
        L.append(f"⏱ Ish vaqti: {s['first']:%H:%M} – {s['last']:%H:%M} ({fmt_min(total)})")
        L.append(f"   Vizitlarda: {fmt_min(s['in_visit'])}{avg} · eng uzoq tanaffus: {fmt_min(s['max_gap'])}")
        if s["short"] or s["open_long"]:
            L.append(f"   ⚠️ 1 daqiqadan qisqa vizit: {s['short']} ta · yopilmay qolgan: {s['open_long']} ta")
    L.append(f"🛒 Buyurtma: <b>{s['orders_n']}</b> ta · {money(s['orders_sum'])} so'm")
    L.append(f"🚚 Yuk chiqqan: <b>{s['ship_n']}</b> ta · {money(s['ship_sum'])} so'm")
    if mon and not mon["plan"]:
        L.append(f"📈 Oy: {short(mon['mtd'])} · prognoz <b>{short(mon['forecast'])}</b> · reja belgilanmagan")
    elif mon:
        L.append(f"📈 Oy: {short(mon['mtd'])} · prognoz <b>{short(mon['forecast'])}</b> / reja {short(mon['plan'])} "
                 f"({mon['pct'] * 100:.0f}%) {status(mon['pct'])}")
        L.append(f"   Reja uchun kuniga {short(mon['need'])} kerak (hozir {short(mon['pace'])}); "
                 f"ziyod uchun {short(mon['need_stretch'])}")
    by = ", ".join(f"{e(k)} {money(v)}" for k, v in s["pay_by"].items())
    L.append(f"💵 Tushgan pul: <b>{money(s['pay_sum'])}</b> so'm" + (f" ({by})" if len(s["pay_by"]) > 0 else ""))
    for c, amt, t in sorted(s["pays"], key=lambda x: -x[1])[:limit]:
        L.append(f"   • {e(c)} — {money(amt)}")
    if len(s["pays"]) > limit:
        L.append(f"   … va yana {len(s['pays']) - limit} ta to'lov")
    L.append(f"↩️ Vozvrat: <b>{s['ret_n']}</b> ta · {money(s['ret_sum'])} so'm")
    for c, summa, reason, prods in s["rets"][:limit]:
        pr = ", ".join(f"{e(n)} ({q:g})" for n, q in prods[:3])
        L.append(f"   • {e(c)} — {money(summa)}" + (f" · {e(reason)}" if reason else "") + (f"\n     {pr}" if pr else ""))
    if s["rejects"]:
        L.append(f"🚫 Otkaz: <b>{len(s['rejects'])}</b> ta")
        for c, r in s["rejects"][:limit]:
            L.append(f"   • {e(str(c))} — {e(r)}")
        if len(s["rejects"]) > limit:
            L.append(f"   … va yana {len(s['rejects']) - limit} ta")
    w = warnings(s)
    if w:
        L.append("⚠️ " + "; ".join(w))
    return "\n".join(L)


def build_messages(S, idle, day, month=None, limit=3900, label=None, title="Agentlar hisoboti", multi=False):
    tot = defaultdict(float)
    for s in S.values():
        tot["planned"] += s["planned"]; tot["visited"] += s["visited"]
        tot["orders_n"] += s["orders_n"]; tot["orders_sum"] += s["orders_sum"]
        tot["pay"] += s["pay_sum"]; tot["ret_n"] += s["ret_n"]; tot["ret_sum"] += s["ret_sum"]
        tot["rej"] += len(s["rejects"])
        tot["ship"] += s["ship_sum"]; tot["ship_n"] += s["ship_n"]
    head = (f"📊 <b>{html.escape(title)} — {label or format(day, '%d.%m.%Y')}</b>\n"
            f"Agentlar: {sum(1 for x in S.values() if x['field'])} ta ishladi"
            + (" + ofis" if any(not x["field"] for x in S.values()) else "") + f" · Vizit: {int(tot['visited'])}"
            + (f" / reja {int(tot['planned'])}" if tot["planned"] else "") + "\n"
            f"🛒 Savdo: <b>{money(tot['orders_sum'])}</b> so'm ({int(tot['orders_n'])} ta buyurtma)\n"
            f"🚚 Yuk chiqqan: <b>{money(tot['ship'])}</b> so'm ({int(tot['ship_n'])} ta)\n"
            f"💵 Tushgan pul: <b>{money(tot['pay'])}</b> so'm\n"
            f"↩️ Vozvrat: {money(tot['ret_sum'])} so'm ({int(tot['ret_n'])} ta)\n"
            f"🚫 Otkaz: {int(tot['rej'])} ta")
    if month:
        t = month["team"]
        head += (f"\n\n📈 <b>Oylik reja:</b> {short(t['plan'])} · oy boshidan {short(t['mtd'])} "
                 f"({t['mtd'] / t['plan'] * 100:.0f}%)\n"
                 f"Prognoz: <b>{short(t['forecast'])}</b> ({t['pct'] * 100:.0f}%) {status(t['pct'])}"
                 f" · oxirgi hafta sur'ati bilan {short(t['forecast_recent'])}\n"
                 f"Rejaga yetish: kuniga <b>{short(t['need_plan'])}</b> kerak (hozir {short(t['pace'])}) · "
                 f"ziyod ({t['stretch_pct']:.0f}%+): kuniga {short(t['need_stretch'])}\n"
                 f"Oy oxirigacha {t['rem']} ish kuni qoldi")
    blocks = [head]
    for s in sorted(S.values(), key=lambda x: -x["orders_sum"]):
        blocks.append(agent_block(s, mon=(month or {}).get("agents", {}).get(s["id"]), multi=multi))
    if idle:
        blocks.append("😴 <b>Hech qanday faoliyat yo'q:</b> " + html.escape(", ".join(idle), quote=False))
    msgs, cur = [], ""
    for b in blocks:
        if cur and len(cur) + len(b) + 2 > limit:
            msgs.append(cur)
            cur = ""
        cur += ("\n\n" if cur else "") + b
    if cur:
        msgs.append(cur)
    return msgs


# ---------------------------------------------------------------- Excel
def make_excel(S, data, day):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Xulosa"
    ws.append(["Agent", "Reja", "Vizit", "GPS", "Foto", "Boshlagan", "Tugatgan", "Ish vaqti (daq)", "Vizitda (daq)",
               "Buyurtma soni", "Savdo (so'm)", "Tushgan pul (so'm)", "Vozvrat soni", "Vozvrat (so'm)", "Otkaz soni", "Yuk soni", "Yuk chiqqan (so'm)"])
    for s in sorted(S.values(), key=lambda x: -x["orders_sum"]):
        tot = round((s["last"] - s["first"]).total_seconds() / 60) if s["first"] else 0
        ws.append([s["name"], s["planned"], s["visited"], s["gps"], s["photo"],
                   f"{s['first']:%H:%M}" if s["first"] else "", f"{s['last']:%H:%M}" if s["last"] else "",
                   tot, round(s["in_visit"]), s["orders_n"], s["orders_sum"], s["pay_sum"],
                   s["ret_n"], s["ret_sum"], len(s["rejects"]), s["ship_n"], s["ship_sum"]])
    w2 = wb.create_sheet("Vizitlar")
    w2.append(["Agent", "Mijoz", "Boshlandi", "Tugadi", "Rejada", "Bo'ldi", "GPS", "Foto", "Buyurtma", "Summa", "Otkaz sababi"])
    names = {a["SD_id"]: a.get("name") for a in data["agents"]}
    for v in data["visits"]:
        w2.append([names.get(v.get("agent_id")) or v.get("agent_name"), v.get("client_name"), v.get("start_date"),
                   v.get("end_date"), v.get("planned"), v.get("visited"), v.get("gps_visit"),
                   v.get("has_photo_report"), v.get("has_order"), v.get("order_summa"), v.get("reject")])
    w3 = wb.create_sheet("Tolovlar")
    w3.append(["Agent", "Mijoz", "Summa", "To'lov turi"])
    w4 = wb.create_sheet("Vozvratlar")
    w4.append(["Agent", "Mijoz", "Summa", "Sabab", "Tovarlar"])
    for s in S.values():
        for c, amt, t in s["pays"]:
            w3.append([s["name"], c, amt, t])
        for c, summa, reason, prods in s["rets"]:
            w4.append([s["name"], c, summa, reason, "; ".join(f"{n} ({q:g})" for n, q in prods)])
    for sheet in wb.worksheets:
        for c in sheet[1]:
            c.font = Font(bold=True)
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 22
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
